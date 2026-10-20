"""Rezyume va vakansiyani professional PDF (reportlab) va Word (python-docx) fayllarga aylantirish.

Dizayn: chapda to'q binafsha yon panel (rasm, aloqa, ko'nikmalar, tillar), o'ngda asosiy qism
(ism, lavozim, o'zi haqida, tajriba, ta'lim, kutilayotgan ish). Rang — web ilovadagi brend rangi #5B57F0.
"""
from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image as PILImage, ImageDraw, ImageOps
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Flowable,
    Frame,
    FrameBreak,
    KeepTogether,
    NextPageTemplate,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.platypus import Image as RLImage
from reportlab.platypus.flowables import HRFlowable

from ..constants import cat_name, edu_name, exp_name, job_type_name, lang_name
from ..utils import age_from_year, lines_list, salary_text

ASSETS = Path(__file__).resolve().parent.parent.parent / "assets"

# ---------- ranglar ----------
PRIMARY = "#5B57F0"
DARK = "#23205C"
DARK_2 = "#34308A"
TEXT = "#12131A"
TEXT_2 = "#4A4E60"
MUTED = "#8C90A3"
LIGHT = "#F3F2FF"
LINE = "#E4E3F5"
SIDE_TEXT = "#E7E6FF"

LBL = {
    "contact": ("ALOQA", "КОНТАКТЫ"),
    "personal": ("SHAXSIY MA'LUMOT", "ЛИЧНЫЕ ДАННЫЕ"),
    "skills": ("KO'NIKMALAR", "НАВЫКИ"),
    "langs": ("TILLAR", "ЯЗЫКИ"),
    "format": ("ISH FORMATI", "ФОРМАТ РАБОТЫ"),
    "about": ("O'ZIM HAQIMDA", "О СЕБЕ"),
    "exp": ("ISH TAJRIBASI", "ОПЫТ РАБОТЫ"),
    "edu": ("TA'LIM", "ОБРАЗОВАНИЕ"),
    "desired": ("KUTILAYOTGAN ISH", "ЖЕЛАЕМАЯ РАБОТА"),
    "age": ("Yoshi", "Возраст"),
    "years": ("yosh", "лет"),
    "born": ("Tug'ilgan yili", "Год рождения"),
    "city": ("Shahar", "Город"),
    "category": ("Soha", "Сфера"),
    "position": ("Lavozim", "Должность"),
    "level": ("Tajriba", "Опыт"),
    "salary": ("Maosh", "Зарплата"),
    "type": ("Format", "Формат"),
    "footer": ("AyolQadami orqali yaratildi", "Создано в AyolQadami"),
    "description": ("ISH HAQIDA", "О РАБОТЕ"),
    "requirements": ("TALABLAR", "ТРЕБОВАНИЯ"),
    "conditions": ("SHAROITLAR", "УСЛОВИЯ"),
    "contacts": ("BOG'LANISH", "КОНТАКТЫ"),
    "vacancy": ("VAKANSIYA", "ВАКАНСИЯ"),
    "no_exp": ("Ish tajribasi ko'rsatilmagan", "Опыт работы не указан"),
}


def L(key: str, lang: str) -> str:
    uz, ru = LBL[key]
    return ru if lang == "ru" else uz


# ---------- shriftlar ----------
def _register_fonts() -> tuple[str, str, str]:
    candidates = [
        (ASSETS / "fonts" / "DejaVuSans.ttf", ASSETS / "fonts" / "DejaVuSans-Bold.ttf", ASSETS / "fonts" / "DejaVuSans-Oblique.ttf"),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
         Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf")),
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf"), Path("C:/Windows/Fonts/ariali.ttf")),
    ]
    for regular, bold, italic in candidates:
        if regular.exists() and bold.exists() and italic.exists():
            pdfmetrics.registerFont(TTFont("AQ", str(regular)))
            pdfmetrics.registerFont(TTFont("AQ-B", str(bold)))
            pdfmetrics.registerFont(TTFont("AQ-I", str(italic)))
            pdfmetrics.registerFontFamily("AQ", normal="AQ", bold="AQ-B", italic="AQ-I", boldItalic="AQ-B")
            return "AQ", "AQ-B", "AQ-I"
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"


F, FB, FI = _register_fonts()


def _x(text) -> str:
    """reportlab Paragraph uchun XML-xavfsiz matn."""
    s = str(text or "")
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def safe_filename(name: str, prefix: str) -> str:
    table = str.maketrans("ʻʼ'`’", "     ")
    slug = re.sub(r"[^A-Za-z0-9]+", "_", (name or "").translate(table)).strip("_")
    return f"{prefix}_{slug or 'AyolQadami'}"


# ---------- rasm ----------
def circle_photo(data: bytes, size: int = 480, ring: int = 14) -> bytes:
    """Rasmni doira shaklida qirqib, oq hoshiya bilan PNG qaytaradi."""
    img = PILImage.open(BytesIO(data)).convert("RGB")
    img = ImageOps.fit(img, (size, size), method=PILImage.LANCZOS, centering=(0.5, 0.35))
    out = PILImage.new("RGBA", (size + ring * 2, size + ring * 2), (0, 0, 0, 0))
    ImageDraw.Draw(out).ellipse((0, 0, size + ring * 2 - 1, size + ring * 2 - 1), fill=(255, 255, 255, 255))
    mask = PILImage.new("L", (size, size), 0)
    ImageDraw.Draw(mask).ellipse((0, 0, size - 1, size - 1), fill=255)
    out.paste(img, (ring, ring), mask)
    buf = BytesIO()
    out.save(buf, "PNG")
    return buf.getvalue()


def initials(name: str) -> str:
    parts = [p for p in re.split(r"\s+", (name or "").strip()) if p]
    return "".join(p[0] for p in parts[:2]).upper() or "AQ"


# ---------- reportlab flowable'lar ----------
class Avatar(Flowable):
    """Rasm bo'lmasa — bosh harflar yozilgan doira."""

    def __init__(self, text: str, size: float):
        super().__init__()
        self.text, self.size = text, size

    def wrap(self, aw, ah):
        self.aw = aw
        return aw, self.size

    def draw(self):
        c, s = self.canv, self.size
        x = (self.aw - s) / 2
        c.setFillColor(colors.white)
        c.circle(x + s / 2, s / 2, s / 2, fill=1, stroke=0)
        c.setFillColor(colors.HexColor(PRIMARY))
        c.circle(x + s / 2, s / 2, s / 2 - 4, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont(FB, s * 0.34)
        c.drawCentredString(x + s / 2, s / 2 - s * 0.12, self.text)


class Chips(Flowable):
    """Ko'nikmalar — yumaloq "tabletka"lar, qatorga sig'maganda keyingi qatorga o'tadi."""

    def __init__(self, items: list[str], fill: str, text_color: str, font_size: float = 8.2):
        super().__init__()
        self.items, self.fill, self.tc, self.fs = items, fill, text_color, font_size
        self.pad_x, self.h, self.gap = 6, font_size + 8, 4

    def _layout(self, aw):
        rows, x, y = [], 0, 0
        for it in self.items:
            text = it if pdfmetrics.stringWidth(it, F, self.fs) <= aw - 2 * self.pad_x else it[:28] + "…"
            w = pdfmetrics.stringWidth(text, F, self.fs) + 2 * self.pad_x
            if x and x + w > aw:
                x, y = 0, y + self.h + self.gap
            rows.append((text, x, y, w))
            x += w + self.gap
        return rows, (y + self.h if rows else 0)

    def wrap(self, aw, ah):
        self.rows, self.height = self._layout(aw)
        return aw, self.height

    def draw(self):
        c = self.canv
        for text, x, y, w in self.rows:
            yy = self.height - y - self.h
            c.setFillColor(colors.HexColor(self.fill))
            c.roundRect(x, yy, w, self.h, self.h / 2, fill=1, stroke=0)
            c.setFillColor(colors.HexColor(self.tc))
            c.setFont(F, self.fs)
            c.drawString(x + self.pad_x, yy + (self.h - self.fs) / 2 + 1.6, text)


def _styles():
    s = {
        "name": ParagraphStyle("name", fontName=FB, fontSize=23, leading=27, textColor=colors.HexColor(TEXT)),
        "headline": ParagraphStyle("headline", fontName=FB, fontSize=12, leading=16, textColor=colors.HexColor(PRIMARY), spaceBefore=2),
        "h": ParagraphStyle("h", fontName=FB, fontSize=10, leading=13, textColor=colors.HexColor(PRIMARY), spaceBefore=14, spaceAfter=3),
        "body": ParagraphStyle("body", fontName=F, fontSize=9.6, leading=14, textColor=colors.HexColor(TEXT_2)),
        "bold": ParagraphStyle("bold", fontName=FB, fontSize=9.8, leading=13.5, textColor=colors.HexColor(TEXT)),
        "muted": ParagraphStyle("muted", fontName=F, fontSize=8.6, leading=12, textColor=colors.HexColor(MUTED)),
        "side_h": ParagraphStyle("side_h", fontName=FB, fontSize=9.2, leading=12, textColor=colors.white, spaceBefore=14, spaceAfter=6),
        "side": ParagraphStyle("side", fontName=F, fontSize=8.8, leading=13, textColor=colors.HexColor(SIDE_TEXT)),
        "side_small": ParagraphStyle("side_small", fontName=F, fontSize=7.8, leading=10, textColor=colors.HexColor("#B9B6F5")),
        "kv_k": ParagraphStyle("kv_k", fontName=F, fontSize=8.6, leading=12, textColor=colors.HexColor(MUTED)),
        "kv_v": ParagraphStyle("kv_v", fontName=FB, fontSize=9.4, leading=12.5, textColor=colors.HexColor(TEXT)),
        "title_w": ParagraphStyle("title_w", fontName=FB, fontSize=20, leading=25, textColor=colors.white),
        "sub_w": ParagraphStyle("sub_w", fontName=F, fontSize=11, leading=15, textColor=colors.HexColor("#E2E1FF")),
        "tag_w": ParagraphStyle("tag_w", fontName=FB, fontSize=8, leading=10, textColor=colors.HexColor("#C9C7FF")),
        "center_w": ParagraphStyle("center_w", fontName=FB, fontSize=8, leading=10, textColor=colors.white, alignment=TA_CENTER),
    }
    return s


def _section(title: str, st, width: float) -> list:
    return [
        Paragraph(_x(title), st["h"]),
        HRFlowable(width=width, thickness=0.8, color=colors.HexColor(LINE), spaceBefore=0, spaceAfter=7),
    ]


EXP_RE = re.compile(r"^\s*([\d.\s/]{4,}\s*[-–—]\s*(?:[\d.\s/]{4,}|hozir\w*|h\.?v\.?|now|present|н\.?\s?в\.?|по\s+н\.?\s?в\.?|наст\w*))\s*[-–—:,]?\s*(.+)$", re.I)


def _exp_rows(text: str | None):
    rows = []
    for ln in lines_list(text):
        m = EXP_RE.match(ln)
        rows.append((m.group(1).strip(), m.group(2).strip(" —–-,")) if m else ("", ln))
    return rows


# ---------- REZYUME PDF ----------
def resume_pdf(r: dict, lang: str = "uz", photo: bytes | None = None) -> bytes:
    buf = BytesIO()
    W, H = A4
    SB = 68 * mm
    st = _styles()
    footer = L("footer", lang)

    def paint(c, doc):
        c.saveState()
        c.setFillColor(colors.HexColor(DARK))
        c.rect(0, 0, SB, H, fill=1, stroke=0)
        # yon panel tepasidagi dekor doiralar (panel chegarasida qirqiladi)
        c.saveState()
        clip = c.beginPath()
        clip.rect(0, 0, SB, H)
        c.clipPath(clip, stroke=0, fill=0)
        c.setFillColor(colors.HexColor(DARK_2))
        c.circle(SB - 6 * mm, H - 4 * mm, 30 * mm, fill=1, stroke=0)
        c.setFillColor(colors.HexColor("#2B2870"))
        c.circle(4 * mm, 30 * mm, 22 * mm, fill=1, stroke=0)
        c.restoreState()
        c.setFillColor(colors.HexColor(PRIMARY))
        c.rect(SB, H - 3.2 * mm, W - SB, 3.2 * mm, fill=1, stroke=0)
        c.setFont(F, 7.4)
        c.setFillColor(colors.HexColor(MUTED))
        c.drawRightString(W - 12 * mm, 9 * mm, f"{footer}  ·  {doc.page}")
        c.setFillColor(colors.HexColor("#8E8AE8"))
        c.drawString(10 * mm, 9 * mm, "AyolQadami")
        c.restoreState()

    left = Frame(9 * mm, 16 * mm, SB - 18 * mm, H - 28 * mm, id="left", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    right_w = W - SB - 22 * mm
    right = Frame(SB + 11 * mm, 16 * mm, right_w, H - 30 * mm, id="right", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    right2 = Frame(SB + 11 * mm, 16 * mm, right_w, H - 30 * mm, id="right2", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc = BaseDocTemplate(buf, pagesize=A4, leftMargin=0, rightMargin=0, topMargin=0, bottomMargin=0,
                          title=f"{r.get('full_name') or ''} — Resume", author="AyolQadami")
    doc.addPageTemplates([PageTemplate("first", [left, right], onPage=paint), PageTemplate("next", [right2], onPage=paint)])

    story: list = [NextPageTemplate("next"), Spacer(1, 6 * mm)]

    # --- yon panel ---
    size = 40 * mm
    if photo:
        try:
            img = RLImage(BytesIO(circle_photo(photo)), width=size, height=size)
            img.hAlign = "CENTER"
            story.append(img)
        except Exception:  # buzuq rasm — bosh harflar
            story.append(Avatar(initials(r.get("full_name")), size))
    else:
        story.append(Avatar(initials(r.get("full_name")), size))
    story.append(Spacer(1, 4 * mm))

    contacts = []
    if r.get("phone"):
        contacts.append(f"☎  {_x(r['phone'])}")
    if r.get("telegram"):
        contacts.append(f"✈  {_x(r['telegram'])}")
    if r.get("city"):
        contacts.append(f"⌂  {_x(r['city'])}")
    if contacts:
        story.append(Paragraph(L("contact", lang), st["side_h"]))
        story += [Paragraph(c, st["side"]) for c in contacts]

    age = age_from_year(r.get("birth_year"))
    if age:
        story.append(Paragraph(L("personal", lang), st["side_h"]))
        story.append(Paragraph(f"{L('age', lang)}: <b>{age} {L('years', lang)}</b>", st["side"]))
        story.append(Paragraph(f"{L('born', lang)}: {r['birth_year']}", st["side"]))

    skills = [s for s in (r.get("skills") or []) if s][:16]
    if skills:
        story.append(Paragraph(L("skills", lang), st["side_h"]))
        story.append(Chips(skills, DARK_2, "#FFFFFF"))

    if r.get("languages"):
        story.append(Paragraph(L("langs", lang), st["side_h"]))
        story += [Paragraph(f"•  {_x(lang_name(x, lang))}", st["side"]) for x in r["languages"][:8]]

    if r.get("job_types"):
        story.append(Paragraph(L("format", lang), st["side_h"]))
        story += [Paragraph(f"•  {_x(job_type_name(x, lang))}", st["side"]) for x in r["job_types"]]

    story.append(FrameBreak())

    # --- asosiy qism ---
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(_x(r.get("full_name")), st["name"]))
    if r.get("headline"):
        story.append(Paragraph(_x(r["headline"]), st["headline"]))
    story.append(HRFlowable(width=18 * mm, thickness=3, color=colors.HexColor(PRIMARY), hAlign="LEFT", spaceBefore=8, spaceAfter=4))

    if r.get("about"):
        story += _section(L("about", lang), st, right_w)
        story.append(Paragraph(_x(r["about"]).replace("\n", "<br/>"), st["body"]))

    story += _section(L("exp", lang), st, right_w)
    if r.get("exp_level"):
        story.append(Paragraph(f"{L('level', lang)}: <b>{_x(exp_name(r['exp_level'], lang))}</b>", st["muted"]))
        story.append(Spacer(1, 4))
    rows = _exp_rows(r.get("experience"))
    if rows:
        data = []
        for period, what in rows:
            data.append([Paragraph(_x(period), st["muted"]), Paragraph(_x(what), st["bold"])])
        has_period = any(p for p, _ in rows)
        widths = [30 * mm, right_w - 30 * mm] if has_period else [0.1, right_w - 0.1]
        tbl = Table(data, colWidths=widths)
        tbl.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (0, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LINEBELOW", (0, 0), (-1, -2), 0.4, colors.HexColor(LINE)),
        ]))
        story.append(tbl)
    elif not r.get("exp_level"):
        story.append(Paragraph(L("no_exp", lang), st["muted"]))

    if r.get("edu_level") or r.get("education"):
        story += _section(L("edu", lang), st, right_w)
        if r.get("edu_level"):
            story.append(Paragraph(_x(edu_name(r["edu_level"], lang)), st["bold"]))
        if r.get("education"):
            story.append(Paragraph(_x(r["education"]).replace("\n", "<br/>"), st["body"]))

    kv = []
    if r.get("headline"):
        kv.append((L("position", lang), r["headline"]))
    if r.get("category"):
        kv.append((L("category", lang), cat_name(r["category"], lang)))
    if r.get("job_types"):
        kv.append((L("type", lang), ", ".join(job_type_name(x, lang) for x in r["job_types"])))
    kv.append((L("salary", lang), salary_text(r.get("salary"), None, lang) if r.get("salary") else ("Договорная" if lang == "ru" else "Kelishiladi")))
    if kv:
        block = _section(L("desired", lang), st, right_w)
        tbl = Table([[Paragraph(_x(k), st["kv_k"]), Paragraph(_x(v), st["kv_v"])] for k, v in kv], colWidths=[34 * mm, right_w - 34 * mm])
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(LIGHT)),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LINEBELOW", (0, 0), (-1, -2), 0.6, colors.white),
        ]))
        block.append(tbl)
        story.append(KeepTogether(block))

    doc.build(story)
    return buf.getvalue()


# ---------- VAKANSIYA PDF ----------
def vacancy_pdf(v: dict, lang: str = "uz") -> bytes:
    buf = BytesIO()
    W, H = A4
    st = _styles()
    margin = 16 * mm
    width = W - 2 * margin
    footer = L("footer", lang)

    def paint(c, doc):
        c.saveState()
        c.setFillColor(colors.HexColor(PRIMARY))
        c.rect(0, H - 3 * mm, W, 3 * mm, fill=1, stroke=0)
        c.setFont(F, 7.4)
        c.setFillColor(colors.HexColor(MUTED))
        c.drawString(margin, 9 * mm, "AyolQadami")
        c.drawRightString(W - margin, 9 * mm, f"{footer}  ·  {doc.page}")
        c.restoreState()

    frame = Frame(margin, 16 * mm, width, H - 30 * mm, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc = BaseDocTemplate(buf, pagesize=A4, title=v.get("title") or "Vacancy", author="AyolQadami")
    doc.addPageTemplates([PageTemplate("all", [frame], onPage=paint)])

    head = [Paragraph(L("vacancy", lang), st["tag_w"]), Spacer(1, 3), Paragraph(_x(v.get("title")), st["title_w"])]
    if v.get("company"):
        head.append(Paragraph(_x(v["company"]), st["sub_w"]))
    head_tbl = Table([[head]], colWidths=[width])
    head_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(PRIMARY)),
        ("LEFTPADDING", (0, 0), (-1, -1), 16), ("RIGHTPADDING", (0, 0), (-1, -1), 16),
        ("TOPPADDING", (0, 0), (-1, -1), 16), ("BOTTOMPADDING", (0, 0), (-1, -1), 16),
        ("ROUNDEDCORNERS", [10, 10, 10, 10]),
    ]))
    story: list = [Spacer(1, 4 * mm), head_tbl, Spacer(1, 5 * mm)]

    info = [(L("salary", lang), salary_text(v.get("salary_from"), v.get("salary_to"), lang))]
    if v.get("city"):
        info.append((L("city", lang), v["city"]))
    if v.get("job_type"):
        info.append((L("type", lang), job_type_name(v["job_type"], lang)))
    if v.get("exp_level"):
        info.append((L("level", lang), exp_name(v["exp_level"], lang)))
    if v.get("category"):
        info.append((L("category", lang), cat_name(v["category"], lang)))
    cells = [[Paragraph(_x(k), st["kv_k"]), Paragraph(_x(val), st["kv_v"])] for k, val in info]
    # 2 ustunli ma'lumot paneli
    grid = []
    for i in range(0, len(cells), 2):
        row = [cells[i]]
        row.append(cells[i + 1] if i + 1 < len(cells) else "")
        grid.append(row)
    info_tbl = Table(grid, colWidths=[width / 2, width / 2])
    info_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(LIGHT)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 12), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LINEBELOW", (0, 0), (-1, -2), 0.8, colors.white), ("LINEAFTER", (0, 0), (0, -1), 0.8, colors.white),
    ]))
    story.append(info_tbl)

    if v.get("description"):
        story += _section(L("description", lang), st, width)
        story.append(Paragraph(_x(v["description"]).replace("\n", "<br/>"), st["body"]))
    for key in ("requirements", "conditions"):
        items = lines_list(v.get(key))
        if items:
            story += _section(L(key, lang), st, width)
            story += [Paragraph(f"<font color='{PRIMARY}'>●</font>&nbsp;&nbsp;{_x(x)}", st["body"]) for x in items]

    contacts = []
    if v.get("contact_phone"):
        contacts.append(f"☎  {_x(v['contact_phone'])}")
    if v.get("contact_tg"):
        contacts.append(f"✈  {_x(v['contact_tg'])}")
    if contacts:
        block = _section(L("contacts", lang), st, width)
        block += [Paragraph(c, st["bold"]) for c in contacts]
        story.append(KeepTogether(block))

    doc.build(story)
    return buf.getvalue()


# ---------- Word yordamchilari ----------
def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tc_pr.append(shd)


def _cell_margins(cell, top=0, bottom=0, left=0, right=0) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(val)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def _no_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _bottom_border(paragraph, color: str = LINE, size: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color.lstrip("#"))
    bdr.append(bottom)
    p_pr.append(bdr)


def _para(container, text="", size=10, bold=False, italic=False, color=TEXT, align=None, before=0, after=2):
    p = container.add_paragraph()
    if text:
        run = p.add_run(str(text))
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = _rgb(color)
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if align:
        p.alignment = align
    return p


def _new_doc(margins_mm: float = 10) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(margins_mm)
    sec.footer_distance = Mm(5)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(2)
    return doc


def _footer(doc: Document, lang: str) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    run = p.add_run(f"AyolQadami  ·  {L('footer', lang)}")
    run.font.size = Pt(7.5)
    run.font.color.rgb = _rgb(MUTED)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


# ---------- REZYUME DOCX ----------
def resume_docx(r: dict, lang: str = "uz", photo: bytes | None = None) -> bytes:
    doc = _new_doc(9)
    _footer(doc, lang)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _no_borders(table)
    row = table.rows[0]
    row.height = Mm(252)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    side, main = row.cells
    side.width, main.width = Mm(64), Mm(128)
    for col, w in zip(table.columns, (Mm(64), Mm(128))):
        col.width = w
    _shade(side, DARK)
    _cell_margins(side, top=360, bottom=240, left=260, right=260)
    _cell_margins(main, top=300, bottom=240, left=420, right=120)

    # --- yon panel ---
    first = side.paragraphs[0]
    first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if photo:
        try:
            first.add_run().add_picture(BytesIO(circle_photo(photo)), width=Mm(40))
        except Exception:
            photo = None
    if not photo:
        run = first.add_run(initials(r.get("full_name")))
        run.font.size, run.font.bold, run.font.color.rgb = Pt(34), True, _rgb("#FFFFFF")

    def side_h(title):
        _para(side, title, size=9.5, bold=True, color="#FFFFFF", before=14, after=4)

    def side_item(text):
        _para(side, text, size=9, color=SIDE_TEXT, after=2)

    contacts = [x for x in (r.get("phone"), r.get("telegram"), r.get("city")) if x]
    if contacts:
        side_h(L("contact", lang))
        for icon, val in zip(("☎", "✈", "⌂"), (r.get("phone"), r.get("telegram"), r.get("city"))):
            if val:
                side_item(f"{icon}  {val}")
    age = age_from_year(r.get("birth_year"))
    if age:
        side_h(L("personal", lang))
        side_item(f"{L('age', lang)}: {age} {L('years', lang)}")
        side_item(f"{L('born', lang)}: {r['birth_year']}")
    if r.get("skills"):
        side_h(L("skills", lang))
        for s in r["skills"][:16]:
            side_item(f"▸  {s}")
    if r.get("languages"):
        side_h(L("langs", lang))
        for x in r["languages"][:8]:
            side_item(f"•  {lang_name(x, lang)}")
    if r.get("job_types"):
        side_h(L("format", lang))
        for x in r["job_types"]:
            side_item(f"•  {job_type_name(x, lang)}")

    # --- asosiy qism ---
    name_p = main.paragraphs[0]
    run = name_p.add_run(r.get("full_name") or "")
    run.font.size, run.font.bold, run.font.color.rgb = Pt(24), True, _rgb(TEXT)
    if r.get("headline"):
        _para(main, r["headline"], size=13, bold=True, color=PRIMARY, after=10)

    def main_h(title):
        p = _para(main, title, size=10.5, bold=True, color=PRIMARY, before=12, after=6)
        _bottom_border(p)

    if r.get("about"):
        main_h(L("about", lang))
        _para(main, r["about"], size=10, color=TEXT_2)

    main_h(L("exp", lang))
    if r.get("exp_level"):
        _para(main, f"{L('level', lang)}: {exp_name(r['exp_level'], lang)}", size=9, color=MUTED, after=4)
    rows = _exp_rows(r.get("experience"))
    for period, what in rows:
        p = main.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        if period:
            pr = p.add_run(period + "   ")
            pr.font.size, pr.font.color.rgb = Pt(9), _rgb(MUTED)
        wr = p.add_run(what)
        wr.font.size, wr.font.bold, wr.font.color.rgb = Pt(10), True, _rgb(TEXT)
    if not rows and not r.get("exp_level"):
        _para(main, L("no_exp", lang), size=9.5, color=MUTED)

    if r.get("edu_level") or r.get("education"):
        main_h(L("edu", lang))
        if r.get("edu_level"):
            _para(main, edu_name(r["edu_level"], lang), size=10, bold=True)
        if r.get("education"):
            _para(main, r["education"], size=10, color=TEXT_2)

    main_h(L("desired", lang))
    kv = []
    if r.get("headline"):
        kv.append((L("position", lang), r["headline"]))
    if r.get("category"):
        kv.append((L("category", lang), cat_name(r["category"], lang)))
    if r.get("job_types"):
        kv.append((L("type", lang), ", ".join(job_type_name(x, lang) for x in r["job_types"])))
    kv.append((L("salary", lang), salary_text(r.get("salary"), None, lang) if r.get("salary") else ("Договорная" if lang == "ru" else "Kelishiladi")))
    kt = main.add_table(rows=0, cols=2)
    kt.autofit = False
    _no_borders(kt)
    for k, val in kv:
        cells = kt.add_row().cells
        cells[0].width, cells[1].width = Mm(36), Mm(84)
        for c in cells:
            _shade(c, LIGHT)
            _cell_margins(c, top=70, bottom=70, left=120, right=80)
        kp = cells[0].paragraphs[0]
        kr = kp.add_run(k)
        kr.font.size, kr.font.color.rgb = Pt(9), _rgb(MUTED)
        vp = cells[1].paragraphs[0]
        vr = vp.add_run(val)
        vr.font.size, vr.font.bold, vr.font.color.rgb = Pt(10), True, _rgb(TEXT)

    # Word jadvaldan keyin majburiy paragraf qo'yadi — u yangi sahifa ochmasligi uchun 1pt qilamiz
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.line_spacing = Pt(1)

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


# ---------- VAKANSIYA DOCX ----------
def vacancy_docx(v: dict, lang: str = "uz") -> bytes:
    doc = _new_doc(16)
    _footer(doc, lang)
    head = doc.add_table(rows=1, cols=1)
    _no_borders(head)
    cell = head.rows[0].cells[0]
    _shade(cell, PRIMARY)
    _cell_margins(cell, top=260, bottom=260, left=300, right=300)
    tag = cell.paragraphs[0].add_run(L("vacancy", lang))
    tag.font.size, tag.font.bold, tag.font.color.rgb = Pt(8), True, _rgb("#C9C7FF")
    _para(cell, v.get("title"), size=20, bold=True, color="#FFFFFF", after=2)
    if v.get("company"):
        _para(cell, v["company"], size=11, color="#E2E1FF", after=0)
    _para(doc, "", after=4)

    info = [(L("salary", lang), salary_text(v.get("salary_from"), v.get("salary_to"), lang))]
    for key, label, fn in (("city", "city", None), ("job_type", "type", job_type_name), ("exp_level", "level", exp_name), ("category", "category", cat_name)):
        if v.get(key):
            info.append((L(label, lang), fn(v[key], lang) if fn else v[key]))
    it = doc.add_table(rows=0, cols=2)
    _no_borders(it)
    it.autofit = False
    for k, val in info:
        cells = it.add_row().cells
        cells[0].width, cells[1].width = Mm(45), Mm(133)
        for c in cells:
            _shade(c, LIGHT)
            _cell_margins(c, top=80, bottom=80, left=160, right=80)
        kr = cells[0].paragraphs[0].add_run(k)
        kr.font.size, kr.font.color.rgb = Pt(9), _rgb(MUTED)
        vr = cells[1].paragraphs[0].add_run(val)
        vr.font.size, vr.font.bold = Pt(10.5), True

    def h(title):
        p = _para(doc, title, size=11, bold=True, color=PRIMARY, before=14, after=6)
        _bottom_border(p)

    if v.get("description"):
        h(L("description", lang))
        _para(doc, v["description"], size=10.5, color=TEXT_2)
    for key in ("requirements", "conditions"):
        items = lines_list(v.get(key))
        if items:
            h(L(key, lang))
            for x in items:
                p = doc.add_paragraph()
                dot = p.add_run("●  ")
                dot.font.color.rgb, dot.font.size = _rgb(PRIMARY), Pt(8)
                p.add_run(x).font.size = Pt(10.5)
    contacts = [x for x in (v.get("contact_phone"), v.get("contact_tg")) if x]
    if contacts:
        h(L("contacts", lang))
        for c in contacts:
            _para(doc, c, size=11, bold=True)

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
